#!/usr/bin/env python3
"""
Téléchargement du corpus WebCognitive (étape 1).

Lit data/sources.json, récupère chaque document dans data/raw/<format>/
et écrit data/corpus_metadata.json (inventaire + empreinte SHA-256).

Usage :
    python scripts/download_corpus.py            # tout télécharger (les fichiers déjà présents sont ignorés)
    python scripts/download_corpus.py --force    # tout re-télécharger
    python scripts/download_corpus.py --only pdf04 docx03
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCES_FILE = ROOT / "data" / "sources.json"
RAW_DIR = ROOT / "data" / "raw"
METADATA_FILE = ROOT / "data" / "corpus_metadata.json"

HEADERS = {
    # Wikipédia exige un User-Agent descriptif.
    "User-Agent": "WebCognitiveCorpus/1.0 (projet etudiant M1 Paris 8; usage pedagogique)",
    "Accept-Language": "fr,en;q=0.8",
}
TIMEOUT = 90
RETRIES = 3

# Sections Wikipédia à ne pas convertir en note Word (bruit pour le futur moteur).
WIKI_STOP_SECTIONS = {
    "notes et références", "références", "notes", "voir aussi", "bibliographie",
    "liens externes", "articles connexes", "annexes",
}
WIKI_NOISE_SELECTORS = [
    "style", "script", "table", "figure", "sup.reference", "sup.mw-ref",
    ".mw-references-wrap", ".reflist", ".navbox", ".bandeau", ".bandeau-container",
    ".homonymie", ".metadata", ".mw-editsection", ".thumb", ".infobox", ".infobox_v2",
    ".infobox_v3", ".hatnote", ".noprint",
]


# --------------------------------------------------------------------------- #
# Réseau
# --------------------------------------------------------------------------- #
def fetch(url: str) -> requests.Response:
    last_error = None
    for attempt in range(1, RETRIES + 1):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=TIMEOUT, allow_redirects=True)
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            last_error = exc
            if attempt < RETRIES:
                time.sleep(2 * attempt)
    raise RuntimeError(f"échec après {RETRIES} tentatives : {last_error}")


def wiki_title(url: str) -> str:
    return unquote(urlparse(url).path.split("/wiki/", 1)[1])


# --------------------------------------------------------------------------- #
# Contrôles de validité (une page anti-robot HTML n'est pas un PDF)
# --------------------------------------------------------------------------- #
def check_content(fmt: str, content: bytes) -> None:
    if fmt == "pdf" and not content.lstrip()[:5].startswith(b"%PDF"):
        raise ValueError("le contenu reçu n'est pas un PDF (page anti-robot ?) : télécharger à la main")
    if fmt == "txt" and b"<html" in content[:500].lower():
        raise ValueError("le contenu reçu est du HTML, pas du texte brut")
    if len(content) < 500:
        raise ValueError(f"fichier suspect, seulement {len(content)} octets")


# --------------------------------------------------------------------------- #
# Conversion Wikipédia -> note Word (.docx)
# --------------------------------------------------------------------------- #
def wiki_html_to_blocks(html: str) -> list[tuple[str, str]]:
    """Transforme le HTML d'un article en blocs (type, texte) : h2/h3/h4, p, li."""
    soup = BeautifulSoup(html, "html.parser")
    for selector in WIKI_NOISE_SELECTORS:
        for node in soup.select(selector):
            node.decompose()

    blocks: list[tuple[str, str]] = []
    for node in soup.find_all(["h2", "h3", "h4", "p", "li"]):
        text = " ".join(node.get_text(" ", strip=True).split())
        if not text:
            continue
        if node.name in ("h2", "h3", "h4"):
            if node.name == "h2" and text.lower() in WIKI_STOP_SECTIONS:
                break
            blocks.append((node.name, text))
        elif node.name == "li":
            # on ignore les listes imbriquées dans d'autres listes (déjà couvertes par le parent)
            if node.find_parent("li") is None and len(text) > 20:
                blocks.append(("li", text))
        else:
            blocks.append(("p", text))
    return blocks


def build_docx(title: str, source_url: str, blocks: list[tuple[str, str]], dest: Path) -> None:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "Wikipédia (contributeurs) - CC BY-SA 4.0"
    doc.core_properties.subject = "Corpus WebCognitive - Histoire d'Internet"

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(f"Note de synthèse : {title}", level=0)
    meta = doc.add_paragraph()
    run = meta.add_run(
        f"Source : {source_url}\n"
        f"Texte issu de Wikipédia, licence CC BY-SA 4.0. "
        f"Récupéré le {datetime.now().strftime('%d/%m/%Y')}."
    )
    run.italic = True
    run.font.size = Pt(9)

    level_map = {"h2": 1, "h3": 2, "h4": 3}
    for kind, text in blocks:
        if kind in level_map:
            doc.add_heading(text, level=level_map[kind])
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(dest)


# --------------------------------------------------------------------------- #
# Téléchargement d'une source
# --------------------------------------------------------------------------- #
def download_source(src: dict, force: bool) -> dict:
    dest = RAW_DIR / src["format"] / src["filename"]
    dest.parent.mkdir(parents=True, exist_ok=True)

    record = {k: src.get(k) for k in ("id", "title", "author", "year", "lang", "type", "format", "url")}
    record["filename"] = src["filename"]
    record["path"] = str(dest.relative_to(ROOT)).replace("\\", "/")

    if dest.exists() and not force:
        record["status"] = "already_present"
    else:
        method = src["method"]
        if method == "direct":
            content = fetch(src["url"]).content
            check_content(src["format"], content)
            dest.write_bytes(content)
        elif method == "wikipedia_html":
            content = fetch(src["url"]).content
            check_content("html", content)
            dest.write_bytes(content)
        elif method == "wikipedia_docx":
            api_url = "https://fr.wikipedia.org/api/rest_v1/page/html/" + requests.utils.quote(
                wiki_title(src["url"]), safe=""
            )
            blocks = wiki_html_to_blocks(fetch(api_url).text)
            if len(blocks) < 5:
                raise ValueError("article vide ou introuvable")
            build_docx(src["title"], src["url"], blocks, dest)
        else:
            raise ValueError(f"méthode inconnue : {method}")
        record["status"] = "downloaded"
        record["retrieved_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")

    data = dest.read_bytes()
    record["size_bytes"] = len(data)
    record["sha256"] = hashlib.sha256(data).hexdigest()
    return record


# --------------------------------------------------------------------------- #
def main() -> int:
    parser = argparse.ArgumentParser(description="Télécharge le corpus WebCognitive.")
    parser.add_argument("--force", action="store_true", help="re-télécharger les fichiers existants")
    parser.add_argument("--only", nargs="*", help="identifiants de sources à traiter (ex. pdf01 txt03)")
    args = parser.parse_args()

    config = json.loads(SOURCES_FILE.read_text(encoding="utf-8"))
    sources = config["sources"]
    if args.only:
        sources = [s for s in sources if s["id"] in set(args.only)]

    previous = {}
    if METADATA_FILE.exists():
        previous = {d["id"]: d for d in json.loads(METADATA_FILE.read_text(encoding="utf-8"))["documents"]}

    failures = []
    for src in sources:
        label = f"[{src['id']}] {src['filename']}"
        try:
            record = download_source(src, args.force)
            # on garde la date de récupération d'origine pour un fichier déjà présent
            if record["status"] == "already_present" and src["id"] in previous:
                record["retrieved_at"] = previous[src["id"]].get("retrieved_at")
            previous[src["id"]] = record
            print(f"OK    {label} ({record['size_bytes'] // 1024} Ko, {record['status']})")
        except Exception as exc:  # noqa: BLE001 - on veut continuer sur les autres sources
            failures.append((src, str(exc)))
            print(f"ECHEC {label} -> {exc}")

    documents = [previous[s["id"]] for s in config["sources"] if s["id"] in previous]
    summary = {}
    for doc in documents:
        summary[doc["format"]] = summary.get(doc["format"], 0) + 1

    METADATA_FILE.write_text(
        json.dumps(
            {
                "theme": config["theme"],
                "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "total_documents": len(documents),
                "by_format": summary,
                "documents": documents,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"\n{len(documents)} documents dans le corpus : {summary}")
    print(f"Inventaire écrit dans {METADATA_FILE.relative_to(ROOT)}")
    if failures:
        print("\nÀ récupérer à la main (ouvrir l'URL dans le navigateur, enregistrer sous le nom indiqué) :")
        for src, err in failures:
            print(f"  - {src['url']}\n    -> data/raw/{src['format']}/{src['filename']}  ({err})")
        print("Puis relancer le script : il ne re-télécharge pas ce qui est déjà là.")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
